from docx import Document

output_path = r'd:\B.Tech\Oral Cancer Detection\Project_Workflow_Oral_Cancer_Detection.docx'

doc = Document()
doc.add_heading('Oral Cancer Detection Project: Workflow and Code Explanation', level=0)

doc.add_paragraph('This project is an AI-powered oral disease detection system designed to analyze tongue images and classify them into common oral health conditions. It combines deep learning, image preprocessing, and a web application to provide predictions with medical explanations.')

doc.add_heading('1. Project Overview', level=1)
doc.add_paragraph('The main purpose of this project is to detect diseases such as oral cancer, hairy tongue, oral lichen, oral thrush, and leukoplakia from uploaded tongue images.')
doc.add_paragraph('The system follows an end-to-end AI workflow: user upload -> image preprocessing -> deep learning prediction -> rule-based disease mapping -> result display.')
doc.add_paragraph('It uses a trained TensorFlow/Keras model, a Flask web backend, and a metadata file containing disease definitions and treatment information.')

doc.add_heading('2. Workflow of the Project', level=1)
workflow_steps = [
    'Step 1: The user opens the web application and uploads a tongue image.',
    'Step 2: The file is sent to the Flask backend through the /index route.',
    'Step 3: The image is read and converted into a usable format using PIL and NumPy.',
    'Step 4: The image is resized to a fixed size and expanded into a batch for model input.',
    'Step 5: The trained model predicts the disease class with the highest probability.',
    'Step 6: The predicted class is matched with the disease metadata stored in details.json.',
    'Step 7: The backend extracts the description, symptoms, causes, and treatment information.',
    'Step 8: The final result is displayed to the user through the browser.'
]
for item in workflow_steps:
    doc.add_paragraph(item)

doc.add_heading('3. Technologies Used', level=1)
for item in [
    'Python: core programming language',
    'Flask: backend web framework',
    'TensorFlow and Keras: deep learning model loading and prediction',
    'NumPy: numerical array processing',
    'Pillow (PIL): image resizing and conversion',
    'HTML and Jinja: frontend rendering',
    'JavaScript: user interactions and preview',
    'JSON: metadata storage for disease details'
]:
    doc.add_paragraph('- ' + item)

doc.add_heading('4. Code Explanation', level=1)

doc.add_paragraph('4.1 File: Code/app.py')
doc.add_paragraph('This file is the main application file. It initializes the Flask app, loads the trained model, defines routes, handles uploaded images, preprocesses them, predicts the disease class, and sends the output to the template.')
doc.add_paragraph('MODEL = tf.keras.models.load_model(r".\\saved_model\\1", compile=False) loads the saved model from the project directory. The model is responsible for image classification.')
doc.add_paragraph('CLASS_NAMES stores all disease categories, such as hairytonguedataset, healthytonguedataset, leokoplakiatonguedataset, oralcancerdataset, orallichensdataset, and oralthrushdataset.')
doc.add_paragraph('The function read_file_as_image() converts the uploaded file into a 256x256 RGB image and then into a NumPy array so it is compatible with the model.')
doc.add_paragraph('The /index route receives the uploaded image, calls MODEL.predict(), determines the highest-probability class, and matches the class with a disease record in the metadata file.')
doc.add_paragraph('The result is then passed to the HTML page through render_template(), where the disease name, description, symptoms, causes, treatments, and confidence are displayed.')

doc.add_paragraph('4.2 File: Code/saved_model/details.json')
doc.add_paragraph('This JSON file stores disease definitions and related information. It links the raw model class label to readable details such as condition name, description, symptoms, causes, and treatment suggestions.')
doc.add_paragraph('This file is important because the model returns only a class label; the JSON file converts that label into understandable medical information for the user.')

doc.add_paragraph('4.3 File: Code/templates/index.html')
doc.add_paragraph('This is the user interface for uploading the tongue image and viewing the model output. It contains the upload form, the image preview area, and sections that display the diagnosis and medical details.')
doc.add_paragraph('The page uses Jinja variables such as prediction, description, symptoms, causes, and treatment to display content dynamically computed by the backend.')

doc.add_paragraph('4.4 File: Code/static/js/main1.js')
doc.add_paragraph('This JavaScript file handles client-side behavior such as image preview, upload, request sending, and status updates. It sends the selected image to the backend and displays the result after the prediction completes.')

doc.add_heading('5. End-to-End Data Flow', level=1)
for item in [
    '1. User selects a tongue image in the browser.',
    '2. Flask receives the uploaded file and reads it.',
    '3. The image is resized and transformed into a format suitable for model input.',
    '4. The trained model predicts the disease category.',
    '5. The predicted label is matched with the JSON medical metadata.',
    '6. The app extracts the disease description, symptoms, causes, and treatment suggestions.',
    '7. The result is displayed on the webpage for the user.'
]:
    doc.add_paragraph(item)

doc.add_heading('6. Importance of the Project', level=1)
for item in [
    'It applies AI to the healthcare domain.',
    'It helps detect oral diseases from images quickly.',
    'It provides support for early diagnosis and awareness.',
    'It combines machine learning, frontend design, and medical data into one system.'
]:
    doc.add_paragraph('- ' + item)

doc.add_heading('7. Conclusion', level=1)
doc.add_paragraph('This project is a complete oral disease detection system built using deep learning and web technologies. It follows a clear workflow: upload image -> preprocess -> predict -> match results with disease metadata -> display diagnosis. It shows how AI can be integrated into a practical healthcare application and make medical screening more automated and accessible.')

doc.save(output_path)
print(f'Created: {output_path}')
